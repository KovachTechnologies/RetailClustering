#!/usr/bin/env python3
"""
Compare basket item intersections for par_inner / par_outer cosine-distance extremes.

Inputs (expected in the current directory or --data-dir):
    par_inner.csv   - same PAR, different external_ids
                      columns: external_id_1, par_1, external_id_2, par_2, cosine_distance
                      first 20 rows = highest cosine distance
                      last  20 rows = lowest  cosine distance
    par_outer.csv   - different PARs, different external_ids
                      same schema and extreme layout
    par_items.csv   - external_id, DSC_ITEM (+ other columns ignored)

Output directory tree (created under --out-dir, default = cwd):

    par_inner/true/     bottom 20 of par_inner  (lowest cosine; SAME par → SHOULD match)
    par_inner/false/    top    20 of par_inner  (highest cosine; SAME par → should match but distance says otherwise)
    par_outer/true/     top    20 of par_outer  (highest cosine; DIFF par → SHOULD NOT match)
    par_outer/false/    bottom 20 of par_outer  (lowest cosine; DIFF par → should not match but distance says otherwise)

Each Excel file is named:
    {ext1}__{ext2}__par{par1}__par{par2}__cd{cosine:.6f}.xlsx

Sheet1 "items":
    All unique DSC_ITEM values from either external_id.
    Intersection items are written in red bold; others in black.
    Columns: item | in_external_1 | in_external_2 | is_intersection

Sheet2 "stats":
    Counts, Jaccard, overlap coefficient, cosine_distance, and directory context.

Usage:
    python compare_par_baskets.py
    python compare_par_baskets.py --data-dir /path/to/csvs --out-dir /path/to/output
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def norm_item(value) -> str:
    """Normalize item description for set operations."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    s = str(value).strip()
    s = re.sub(r"\s+", " ", s)
    return s.casefold()


def sanitize_filename(s: str) -> str:
    """Make a string safe for use in a filename."""
    s = re.sub(r"[^\w.\-]+", "_", str(s))
    return s.strip("_")[:80] or "unknown"


def load_items(path: Path) -> dict[str, set[str]]:
    """external_id -> set of normalized DSC_ITEM."""
    df = pd.read_csv(path, usecols=lambda c: c in ("external_id", "DSC_ITEM") or
                     c.lower() in ("external_id", "dsc_item"),
                     dtype=str)
    # flexible column matching
    cols = {c.lower(): c for c in df.columns}
    id_col = cols.get("external_id")
    item_col = cols.get("dsc_item")
    if not id_col or not item_col:
        raise ValueError(f"par_items.csv must contain external_id and DSC_ITEM columns; found {list(df.columns)}")

    grouped: dict[str, set[str]] = defaultdict(set)
    for _, row in df.iterrows():
        eid = str(row[id_col]).strip() if pd.notna(row[id_col]) else ""
        item = norm_item(row[item_col])
        if eid and item:
            grouped[eid].add(item)
    return dict(grouped)


def load_pairs(path: Path) -> pd.DataFrame:
    """Load a par_inner or par_outer CSV and coerce types."""
    df = pd.read_csv(path, dtype=str)
    # normalize expected column names (case-insensitive)
    rename = {}
    lower = {c.lower(): c for c in df.columns}
    for want in ("external_id_1", "par_1", "external_id_2", "par_2", "cosine_distance"):
        if want in lower:
            rename[lower[want]] = want
        else:
            # try common variants
            for c in df.columns:
                if want.replace("_", "") in c.lower().replace("_", ""):
                    rename[c] = want
                    break
    df = df.rename(columns=rename)
    required = {"external_id_1", "par_1", "external_id_2", "par_2", "cosine_distance"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"{path.name} missing columns: {missing}. Found: {list(df.columns)}")

    df["cosine_distance"] = pd.to_numeric(df["cosine_distance"], errors="coerce")
    df = df.dropna(subset=["cosine_distance"]).reset_index(drop=True)
    return df


def write_excel(
    out_path: Path,
    ext1: str,
    ext2: str,
    par1: str,
    par2: str,
    cosine: float,
    items1: set[str],
    items2: set[str],
    context: str,
) -> None:
    """Write one Excel workbook for a single pair."""
    intersection = items1 & items2
    only1 = items1 - items2
    only2 = items2 - items1
    union = items1 | items2

    # sort for stable output: intersection first, then only1, then only2
    ordered = (
        sorted(intersection)
        + sorted(only1)
        + sorted(only2)
    )

    wb = Workbook()

    # ---------- Sheet 1: items ----------
    ws = wb.active
    ws.title = "items"

    red_font = Font(color="FF0000", bold=True)
    black_font = Font(color="000000")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    alt_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    headers = ["item", "in_external_1", "in_external_2", "is_intersection"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(ordered, start=2):
        is_inter = item in intersection
        row = [
            item,
            "Y" if item in items1 else "",
            "Y" if item in items2 else "",
            "Y" if is_inter else "",
        ]
        ws.append(row)
        font = red_font if is_inter else black_font
        for col in range(1, 5):
            cell = ws.cell(row=i, column=col)
            cell.font = font
            cell.border = thin
            if i % 2 == 0 and not is_inter:
                cell.fill = alt_fill

    # column widths
    widths = [60, 18, 18, 18]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    ws.freeze_panes = "A2"

    # ---------- Sheet 2: stats ----------
    stats = wb.create_sheet("stats")
    stats.append(["metric", "value"])
    for cell in stats[1]:
        cell.font = header_font
        cell.fill = header_fill

    jaccard = len(intersection) / len(union) if union else 0.0
    overlap = len(intersection) / min(len(items1), len(items2)) if items1 and items2 else 0.0

    rows = [
        ("context", context),
        ("external_id_1", ext1),
        ("external_id_2", ext2),
        ("par_1", par1),
        ("par_2", par2),
        ("same_par", "Y" if str(par1) == str(par2) else "N"),
        ("cosine_distance", round(float(cosine), 6)),
        ("n_items_external_1", len(items1)),
        ("n_items_external_2", len(items2)),
        ("n_intersection", len(intersection)),
        ("n_only_1", len(only1)),
        ("n_only_2", len(only2)),
        ("n_union", len(union)),
        ("jaccard", round(jaccard, 4)),
        ("overlap_coefficient", round(overlap, 4)),
    ]
    for r in rows:
        stats.append(list(r))

    stats.column_dimensions["A"].width = 28
    stats.column_dimensions["B"].width = 40
    for row in stats.iter_rows(min_row=2, max_row=stats.max_row, max_col=2):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center")

    # freeze and nice header
    stats.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def process_file(
    pairs_df: pd.DataFrame,
    items: dict[str, set[str]],
    base_dir: Path,
    is_inner: bool,
) -> None:
    """Process one of the pair files (inner or outer) and emit the four directories."""
    n = len(pairs_df)
    if n < 40:
        print(f"Warning: {n} rows found; expected at least 40 for top/bottom 20.")

    top20 = pairs_df.head(20)      # highest cosine distance
    bot20 = pairs_df.tail(20)      # lowest cosine distance

    if is_inner:
        # par_1 == par_2 expected
        # true  = bottom 20 (low cosine, SHOULD match)
        # false = top    20 (high cosine, should match but distance says no)
        mapping = [
            (bot20, base_dir / "par_inner" / "true",  "par_inner/true  (lowest cosine, same PAR → expect match)"),
            (top20, base_dir / "par_inner" / "false", "par_inner/false (highest cosine, same PAR → expect match but high distance)"),
        ]
    else:
        # par_1 != par_2 expected
        # true  = top    20 (high cosine, SHOULD NOT match)
        # false = bottom 20 (low cosine, should not match but low distance)
        mapping = [
            (top20, base_dir / "par_outer" / "true",  "par_outer/true  (highest cosine, different PAR → expect no match)"),
            (bot20, base_dir / "par_outer" / "false", "par_outer/false (lowest cosine, different PAR → expect no match but low distance)"),
        ]

    for subset, out_dir, context in mapping:
        out_dir.mkdir(parents=True, exist_ok=True)
        print(f"\nWriting {len(subset)} files to {out_dir}")
        for _, row in subset.iterrows():
            ext1 = str(row["external_id_1"]).strip()
            ext2 = str(row["external_id_2"]).strip()
            par1 = str(row["par_1"]).strip()
            par2 = str(row["par_2"]).strip()
            cosine = float(row["cosine_distance"])

            items1 = items.get(ext1, set())
            items2 = items.get(ext2, set())

            fname = (
                f"{sanitize_filename(ext1)}__{sanitize_filename(ext2)}"
                f"__par{sanitize_filename(par1)}__par{sanitize_filename(par2)}"
                f"__cd{cosine:.6f}.xlsx"
            )
            out_path = out_dir / fname
            write_excel(out_path, ext1, ext2, par1, par2, cosine,
                        items1, items2, context)
            print(f"  {fname}  (|∩|={len(items1 & items2)}, |1|={len(items1)}, |2|={len(items2)})")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data-dir", type=Path, default=Path("."),
                        help="Directory containing par_inner.csv, par_outer.csv, par_items.csv")
    parser.add_argument("--out-dir", type=Path, default=Path("."),
                        help="Root directory under which par_inner/ and par_outer/ will be created")
    args = parser.parse_args()

    data_dir = args.data_dir.resolve()
    out_dir = args.out_dir.resolve()

    inner_path = data_dir / "par_inner.csv"
    outer_path = data_dir / "par_outer.csv"
    items_path = data_dir / "par_items.csv"

    for p in (inner_path, outer_path, items_path):
        if not p.exists():
            raise SystemExit(f"Missing required file: {p}")

    print("Loading par_items.csv …")
    items = load_items(items_path)
    print(f"  {len(items)} distinct external_ids with items")

    print("Loading par_inner.csv …")
    inner_df = load_pairs(inner_path)
    print(f"  {len(inner_df)} rows")

    print("Loading par_outer.csv …")
    outer_df = load_pairs(outer_path)
    print(f"  {len(outer_df)} rows")

    print("\n=== Processing par_inner ===")
    process_file(inner_df, items, out_dir, is_inner=True)

    print("\n=== Processing par_outer ===")
    process_file(outer_df, items, out_dir, is_inner=False)

    print("\nDone.")
    print(f"Output written under:\n  {out_dir / 'par_inner'}\n  {out_dir / 'par_outer'}")


if __name__ == "__main__":
    main()
